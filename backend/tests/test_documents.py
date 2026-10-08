"""Tests for the PDF and Word versions of the deliverables."""

import io

import pytest
from docx import Document as WordDocument
from pypdf import PdfReader

from app.pipeline import documents
from app.pipeline.documents import build_documents, render_docx, render_pdf
from app.pipeline.export import build_exports
from app.schemas import (
    ActionItem,
    Decision,
    MeetingMinute,
    MeetingRecord,
    Proposal,
    RawTranscript,
    RefinedTranscript,
    RefinementStatus,
    RejectedItem,
    Segment,
    SegmentRefinement,
)

RAW_TEXTS = ["we agreed to use cube nettees", "priya will update the docs by friday", "send the cost report"]
REFINED_TEXTS = ["We agreed to use Kubernetes.", "Priya will update the docs by Friday.", "Send the cost report."]


def make():
    raw = RawTranscript(segments=[Segment(id=i, start=i * 5.0, end=i * 5.0 + 4, text=t)
                                  for i, t in enumerate(RAW_TEXTS)], stt_model="whisper-large-v3",
                        duration_seconds=15)
    refs = [SegmentRefinement(segment_id=i, start=i * 5.0, end=i * 5.0 + 4, original_text=r, refined_text=f,
                              changed=True, status=RefinementStatus.REFINED)
            for i, (r, f) in enumerate(zip(RAW_TEXTS, REFINED_TEXTS))]
    refined = RefinedTranscript(segments=[Segment(id=i, start=i * 5.0, end=i * 5.0 + 4, text=t)
                                          for i, t in enumerate(REFINED_TEXTS)],
                                refinements=refs, refine_model="gemini-3.6-flash")
    record = MeetingRecord(
        summary="The team agreed to use Kubernetes.",
        minutes=[MeetingMinute(topic="Platform", discussion="Kubernetes was chosen.", segment_ids=[0])],
        decisions=[Decision(decision="Use Kubernetes", evidence_segment_ids=[0],
                            evidence_quote="We agreed to use Kubernetes.")],
        proposals=[Proposal(proposal="Try serverless later", evidence_segment_ids=[0], evidence_quote="use")],
        action_items=[
            ActionItem(task="Update the docs", owner="Priya", deadline="Friday", evidence_segment_ids=[1],
                       evidence_quote="Priya will update the docs by Friday."),
            ActionItem(task="Send the cost report", evidence_segment_ids=[2],
                       evidence_quote="Send the cost report."),
        ],
        rejected_items=[RejectedItem(kind="action_item", content={"task": "Invented task"},
                                     reasons=["owner invented"])],
        document_model="gemini-3-flash-preview",
    )
    return raw, refined, record


def pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return " ".join(" ".join((page.extract_text() or "") for page in reader.pages).split())


def docx_text(data: bytes) -> str:
    doc = WordDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    parts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return " ".join(" ".join(parts).split())


def rendered(doc):
    return pdf_text(render_pdf(doc)), docx_text(render_docx(doc))


@pytest.fixture
def docs():
    raw, refined, record = make()
    return {d.key: d for d in build_documents(raw, refined, record, source_name="meeting.wav")}


def test_all_six_deliverables_exist(docs):
    assert list(docs) == ["raw_transcript", "refined_transcript", "minutes", "decisions", "action_items",
                          "meeting_record"]


def test_files_are_valid_pdf_and_docx(docs):
    for doc in docs.values():
        pdf = render_pdf(doc)
        assert pdf.startswith(b"%PDF-") and len(PdfReader(io.BytesIO(pdf)).pages) >= 1
        assert WordDocument(io.BytesIO(render_docx(doc))).paragraphs


def test_raw_transcript_is_verbatim(docs):
    pdf, word = rendered(docs["raw_transcript"])
    for text in (pdf, word):
        assert "Raw transcript" in text and "meeting.wav" in text and "whisper-large-v3" in text
        assert "we agreed to use cube nettees" in text and "00:00:05" in text


def test_refined_transcript_lists_corrections(docs):
    pdf, word = rendered(docs["refined_transcript"])
    for text in (pdf, word):
        assert "We agreed to use Kubernetes." in text
        assert "Corrections made" in text and "cube nettees" in text  # raw shown next to the correction


def test_minutes(docs):
    pdf, word = rendered(docs["minutes"])
    for text in (pdf, word):
        assert "Meeting minutes" in text and "The team agreed to use Kubernetes." in text
        assert "Platform" in text and "Kubernetes was chosen." in text


def test_key_decisions_keep_proposals_separate(docs):
    pdf, word = rendered(docs["decisions"])
    for text in (pdf, word):
        assert "Use Kubernetes" in text and "We agreed to use Kubernetes." in text
        assert "Discussed but not agreed" in text
        assert text.index("Use Kubernetes") < text.index("Discussed but not agreed") < text.index("Try serverless")


def test_action_items_show_not_specified_for_missing_values(docs):
    pdf, word = rendered(docs["action_items"])
    for text in (pdf, word):
        assert "Update the docs" in text and "Priya" in text and "Friday" in text
        assert "Send the cost report" in text and text.count("Not specified") >= 2
        assert "Invented task" not in text, "withheld items are not presented as tasks"


def test_complete_record_has_everything(docs):
    pdf, word = rendered(docs["meeting_record"])
    for text in (pdf, word):
        for expected in ["Summary", "Key decisions", "Proposals (not agreed)", "Action items",
                         "Withheld by the safety checks", "owner invented", "Appendix A", "Appendix B",
                         "gemini-3.6-flash", "gemini-3-flash-preview", "whisper-large-v3"]:
            assert expected in text, expected


def test_empty_record_says_so_explicitly():
    raw, refined, _ = make()
    record = MeetingRecord(summary="Nothing was decided.", document_model="m")
    docs = {d.key: d for d in build_documents(raw, refined, record)}
    for text in rendered(docs["decisions"]):
        assert "No decisions were reached in this meeting." in text and "No open proposals." in text
    for text in rendered(docs["action_items"]):
        assert "No action items were assigned in this meeting." in text


def test_partial_results_only_produce_available_documents():
    raw, _, _ = make()
    assert [d.key for d in build_documents(raw, None, None)] == ["raw_transcript"]
    assert build_documents(None, None, None) == []


def test_pdf_and_word_match_the_json_export():
    raw, refined, record = make()
    files = {f.key: f for f in build_exports(raw, refined, record)}
    for d in record.decisions:
        assert d.decision in pdf_text(files["decisions_pdf"].data)
        assert d.decision in docx_text(files["decisions_docx"].data)
    for a in record.action_items:
        assert a.task in pdf_text(files["action_items_pdf"].data)
        assert a.task in docx_text(files["action_items_docx"].data)


def test_unicode_is_preserved_with_a_unicode_font():
    if documents.find_unicode_font() is None:
        pytest.skip("no Unicode TTF font installed")
    raw, refined, record = make()
    record.action_items[1].task = "Approve ₹50,000 → next sprint"
    doc = {d.key: d for d in build_documents(raw, refined, record)}["action_items"]
    pdf, word = rendered(doc)
    assert "₹50,000 → next sprint" in word
    assert "50,000" in pdf and "next sprint" in pdf


def test_pdf_falls_back_without_a_unicode_font(monkeypatch):
    monkeypatch.setattr(documents, "find_unicode_font", lambda: None)
    raw, refined, record = make()
    record.action_items[1].task = "Approve ₹50,000 → “next” sprint"
    doc = {d.key: d for d in build_documents(raw, refined, record)}["action_items"]
    text = pdf_text(render_pdf(doc))
    assert 'Approve Rs.50,000 -> "next" sprint' in text


def test_control_characters_do_not_break_word_files():
    raw, refined, record = make()
    record.summary = "Bad\x00 char\x0b here"
    doc = {d.key: d for d in build_documents(raw, refined, record)}["minutes"]
    assert "Bad char here" in docx_text(render_docx(doc))
