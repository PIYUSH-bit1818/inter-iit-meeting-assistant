"""PDF and Word (.docx) versions of the meeting deliverables.

Each deliverable (raw transcript, refined transcript, meeting minutes, key
decisions, action items, complete record) is first described as a list of
neutral blocks - title, metadata, headings, paragraphs, quotes, tables,
transcript lines - built from the same RawTranscript / RefinedTranscript /
MeetingRecord objects as every other export. The same blocks are then
rendered to PDF (fpdf2) and DOCX (python-docx), so both formats always carry
identical content. Missing owners and deadlines appear as "Not specified".
"""

from __future__ import annotations

import io
import logging
import os
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from ..schemas import MeetingRecord, RawTranscript, RefinedTranscript, Segment

BRAND = "Scripted"
TAGLINE = "AI meeting assistant"
NOT_SPECIFIED = "Not specified"
PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# ---------------------------------------------------------------------------
# Neutral document blocks
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Title:
    text: str
    subtitle: str = ""


@dataclass(frozen=True)
class Meta:
    items: tuple[tuple[str, str], ...]


@dataclass(frozen=True)
class Heading:
    text: str
    level: int = 1


@dataclass(frozen=True)
class Para:
    text: str
    muted: bool = False


@dataclass(frozen=True)
class Quote:
    text: str
    ref: str = ""


@dataclass(frozen=True)
class Table:
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    widths: tuple[float, ...]  # relative column widths


@dataclass(frozen=True)
class Lines:
    """Transcript lines: (time/segment label, text)."""

    rows: tuple[tuple[str, str], ...]


Block = Title | Meta | Heading | Para | Quote | Table | Lines


@dataclass(frozen=True)
class Document:
    key: str  # base export key, e.g. "minutes"
    title: str
    filename_stem: str
    blocks: tuple[Block, ...]


# ---------------------------------------------------------------------------
# Building the deliverables
# ---------------------------------------------------------------------------


def _ts(seconds: float | None) -> str:
    total = int(seconds or 0)
    return f"{total // 3600:02d}:{total % 3600 // 60:02d}:{total % 60:02d}"


def _refs(ids: list[int], segments: dict[int, Segment]) -> str:
    return ", ".join(f"#{i:02d} ({_ts(segments[i].start)})" if i in segments else f"#{i:02d}" for i in ids)


def _lines(segments: list[Segment]) -> Lines:
    return Lines(tuple((f"{_ts(s.start)}  #{s.id:02d}", s.text.strip()) for s in segments if s.text.strip()))


def _meta(source: str | None, *pairs: tuple[str, object]) -> Meta:
    items = [("Recording", source)] if source else []
    items += [(k, str(v)) for k, v in pairs if v not in (None, "")]
    return Meta(tuple(items))


def _decision_blocks(record: MeetingRecord, segs: dict[int, Segment]) -> list[Block]:
    if not record.decisions:
        return [Para("No decisions were reached in this meeting.")]
    blocks: list[Block] = []
    for n, d in enumerate(record.decisions, 1):
        blocks += [Heading(f"{n}. {d.decision}", 2),
                   Quote(d.evidence_quote, "Evidence: segments " + _refs(d.evidence_segment_ids, segs))]
    return blocks


def _proposal_blocks(record: MeetingRecord, segs: dict[int, Segment]) -> list[Block]:
    blocks: list[Block] = [Para("Raised during the meeting but not confirmed. These are not decisions.", muted=True)]
    if not record.proposals:
        return blocks + [Para("No open proposals.")]
    for n, p in enumerate(record.proposals, 1):
        blocks += [Heading(f"{n}. {p.proposal}", 2),
                   Quote(p.evidence_quote, "Evidence: segments " + _refs(p.evidence_segment_ids, segs))]
    return blocks


def _action_blocks(record: MeetingRecord, segs: dict[int, Segment]) -> list[Block]:
    note = Para("Owner and deadline are shown only when the recording states them; otherwise "
                f"they are marked \u201c{NOT_SPECIFIED}\u201d.", muted=True)
    if not record.action_items:
        return [Para("No action items were assigned in this meeting."), note]
    rows = tuple(
        (str(n), a.task, a.owner or NOT_SPECIFIED, a.deadline or NOT_SPECIFIED,
         f"\u201c{a.evidence_quote}\u201d \u2014 segments {_refs(a.evidence_segment_ids, segs)}")
        for n, a in enumerate(record.action_items, 1))
    return [Table(("#", "Task", "Owner", "Deadline", "Evidence"), rows, (5, 27, 14, 16, 38)), note]


def _minutes_blocks(record: MeetingRecord, segs: dict[int, Segment]) -> list[Block]:
    blocks: list[Block] = [Heading("Summary"),
                           Para(record.summary or "The summary was withheld because it failed validation.")]
    blocks.append(Heading("Discussion points"))
    if not record.minutes:
        return blocks + [Para("No minutes were recorded.")]
    for m in record.minutes:
        blocks += [Heading(m.topic, 2), Para(m.discussion), Para("Segments " + _refs(m.segment_ids, segs), muted=True)]
    return blocks


def _models(raw, refined, record) -> list[tuple[str, object]]:
    return [("Speech-to-text model", raw.stt_model if raw else None),
            ("Refinement model", refined.refine_model if refined else None),
            ("Documentation model", record.document_model if record else None)]


def build_documents(
    raw: RawTranscript | None,
    refined: RefinedTranscript | None,
    record: MeetingRecord | None,
    source_name: str | None = None,
) -> list[Document]:
    """The deliverables available for a meeting (partial results allowed)."""
    docs: list[Document] = []
    segs = {s.id: s for s in (refined.segments if refined else raw.segments if raw else [])}
    duration = _ts(raw.duration_seconds) if raw and raw.duration_seconds else None

    if raw:
        docs.append(Document("raw_transcript", "Raw transcript", "raw_transcript", (
            Title("Raw transcript", "Speech-to-text result before language-model refinement"),
            _meta(source_name, ("Speech-to-text model", raw.stt_model), ("Duration", duration),
                  ("Segments", len(raw.segments))),
            Para("Exactly as returned by the speech-to-text model. It is never modified.", muted=True),
            _lines(raw.segments),
        )))

    if refined:
        refs = refined.refinements
        corrected = [r for r in refs if r.status.value == "refined"]
        rejected = [r for r in refs if r.status.value == "fallback"]
        blocks: list[Block] = [
            Title("Refined transcript", "Transcript after domain-aware terminology correction"),
            _meta(source_name, ("Refinement model", refined.refine_model), ("Duration", duration),
                  ("Segments", len(refined.segments)), ("Segments corrected", len(corrected)),
                  ("Corrections rejected", len(rejected))),
            Para("Same segments and timestamps as the raw transcript. Every correction was checked; any edit that "
                 "could change the meaning was rejected and the raw text kept.", muted=True),
            _lines(refined.segments),
            Heading("Corrections made"),
        ]
        if corrected:
            blocks.append(Table(("Segment", "Raw (speech-to-text)", "Refined"),
                                tuple((f"#{r.segment_id:02d} {_ts(r.start)}", r.original_text, r.refined_text)
                                      for r in corrected), (14, 43, 43)))
        else:
            blocks.append(Para("No corrections were needed."))
        if rejected:
            blocks += [Heading("Corrections rejected by the safety checks"),
                       Table(("Segment", "Rejected edit", "Reason"),
                             tuple((f"#{r.segment_id:02d} {_ts(r.start)}", r.rejected_text or "", "; ".join(r.issues))
                                   for r in rejected), (14, 50, 36))]
        docs.append(Document("refined_transcript", "Refined transcript", "refined_transcript", tuple(blocks)))

    if record:
        meta = _meta(source_name, ("Duration", duration), ("Documentation model", record.document_model))
        docs.append(Document("minutes", "Meeting minutes", "meeting_minutes", (
            Title("Meeting minutes", "Concise summary and organized account of the main discussion points"),
            meta, *_minutes_blocks(record, segs))))
        docs.append(Document("decisions", "Key decisions", "key_decisions", (
            Title("Key decisions", "Decisions explicitly reached in the meeting, with their evidence"),
            meta, Heading("Decisions"), *_decision_blocks(record, segs),
            Heading("Discussed but not agreed"), *_proposal_blocks(record, segs))))
        docs.append(Document("action_items", "Action items", "action_items", (
            Title("Action items", "Assigned tasks with any stated owner or deadline"),
            meta, *_action_blocks(record, segs))))

        full: list[Block] = [
            Title("Meeting record", f"Complete record generated by {BRAND}"),
            _meta(source_name, ("Duration", duration), *_models(raw, refined, record)),
            *_minutes_blocks(record, segs),
            Heading("Key decisions"), *_decision_blocks(record, segs),
            Heading("Proposals (not agreed)"), *_proposal_blocks(record, segs),
            Heading("Action items"), *_action_blocks(record, segs),
        ]
        if record.rejected_items:
            full += [Heading("Withheld by the safety checks"),
                     Para("Generated items that could not be grounded in the transcript. They are not part of "
                          "the record.", muted=True),
                     Table(("Type", "Attempted content", "Why it was withheld"),
                           tuple((r.kind.value.replace("_", " "), _attempted(r.content), "; ".join(r.reasons))
                                 for r in record.rejected_items), (14, 46, 40))]
        if refined:
            full += [Heading("Appendix A \u2014 Refined transcript"), _lines(refined.segments)]
        if raw:
            full += [Heading("Appendix B \u2014 Raw transcript"), _lines(raw.segments)]
        docs.append(Document("meeting_record", "Complete meeting record", "meeting_record", tuple(full)))
    return docs


def _attempted(content: dict) -> str:
    for key in ("decision", "proposal", "task", "discussion", "summary"):
        if content.get(key):
            return str(content[key])
    return ""


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# Used only when no Unicode font is available for the PDF
_LATIN1 = str.maketrans({"\u201c": '"', "\u201d": '"', "\u2018": "'", "\u2019": "'", "\u2013": "-",
                         "\u2014": "-", "\u2192": "->", "\u2026": "...", "\u2022": "*", "\u20b9": "Rs.",
                         "\u20ac": "EUR ", "\u00a0": " "})


def _clean(text: str) -> str:
    return _CONTROL.sub("", text or "")


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_FONT_CANDIDATES = (
    ("{windir}/Fonts/segoeui.ttf", "{windir}/Fonts/segoeuib.ttf"),
    ("{windir}/Fonts/arial.ttf", "{windir}/Fonts/arialbd.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
)


@lru_cache
def find_unicode_font() -> tuple[str, str] | None:
    """A (regular, bold) TTF pair with broad Unicode coverage, if one is installed.

    PDF_FONT_REGULAR / PDF_FONT_BOLD can point at specific files.
    """
    env = (os.environ.get("PDF_FONT_REGULAR"), os.environ.get("PDF_FONT_BOLD"))
    if all(env) and all(Path(p).is_file() for p in env):
        return env[0], env[1]
    windir = os.environ.get("WINDIR", r"C:\Windows")
    for regular, bold in _FONT_CANDIDATES:
        pair = (regular.format(windir=windir), bold.format(windir=windir))
        if all(Path(p).is_file() for p in pair):
            return pair
    return None


_NAVY = (11, 31, 69)
_BLUE = (29, 78, 216)
_TEXT = (30, 41, 59)
_MUTED = (100, 116, 139)
_RULE = (203, 213, 225)
_ZEBRA = (244, 247, 252)


def render_pdf(doc: Document) -> bytes:
    from fpdf import FPDF
    from fpdf.enums import TableCellFillMode, XPos, YPos
    from fpdf.fonts import FontFace

    # fontTools warns about font tables it cannot subset (harmless; they are dropped)
    logging.getLogger("fontTools.subset").setLevel(logging.ERROR)
    font = find_unicode_font()
    family = "Body" if font else "Helvetica"
    text = (lambda s: _clean(s)) if font else (lambda s: _clean(s).translate(_LATIN1)
                                                 .encode("latin-1", "replace").decode("latin-1"))

    class _PDF(FPDF):
        def header(self) -> None:
            self.set_y(9)
            self.set_font(family, "B", 8.5)
            self.set_text_color(*_BLUE)
            self.cell(self.epw / 2, 5, f"{BRAND.upper()}  \u00b7  {TAGLINE}" if font else f"{BRAND.upper()} - {TAGLINE}",
                      new_x=XPos.RIGHT, new_y=YPos.TOP)
            self.set_font(family, "", 8.5)
            self.set_text_color(*_MUTED)
            self.cell(self.epw / 2, 5, text(doc.title), align="R", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            self.set_draw_color(*_RULE)
            self.set_line_width(0.2)
            self.line(self.l_margin, self.get_y() + 1.5, self.w - self.r_margin, self.get_y() + 1.5)
            self.ln(7)

        def footer(self) -> None:
            self.set_y(-13)
            self.set_font(family, "", 8)
            self.set_text_color(*_MUTED)
            self.cell(0, 6, f"Generated by {BRAND}  \u00b7  Page {self.page_no()} of {{nb}}" if font
                      else f"Generated by {BRAND} - Page {self.page_no()} of {{nb}}", align="C")

    pdf = _PDF(format="A4")
    if font:
        pdf.add_font(family, "", font[0])
        pdf.add_font(family, "B", font[1])
    pdf.set_margins(18, 16, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.set_title(f"{doc.title} - {BRAND}")
    pdf.set_author(BRAND)
    pdf.add_page()

    def para(value: str, size: float = 10.5, color=_TEXT, style: str = "", h: float = 5.4) -> None:
        pdf.set_font(family, style, size)
        pdf.set_text_color(*color)
        pdf.multi_cell(0, h, text(value), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    for block in doc.blocks:
        if isinstance(block, Title):
            para(block.text, 21, _NAVY, "B", 10)
            if block.subtitle:
                para(block.subtitle, 11, _MUTED, h=6)
            pdf.ln(2)
        elif isinstance(block, Meta):
            pdf.set_fill_color(*_ZEBRA)
            for label, value in block.items:
                pdf.set_font(family, "B", 9)
                pdf.set_text_color(*_MUTED)
                pdf.cell(44, 5.6, text(label.upper()), new_x=XPos.RIGHT, new_y=YPos.TOP)
                pdf.set_font(family, "", 9.5)
                pdf.set_text_color(*_TEXT)
                pdf.multi_cell(0, 5.6, text(value), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            pdf.set_draw_color(*_RULE)
            pdf.line(pdf.l_margin, pdf.get_y() + 2, pdf.w - pdf.r_margin, pdf.get_y() + 2)
            pdf.ln(6)
        elif isinstance(block, Heading):
            pdf.ln(3 if block.level == 1 else 1.5)
            if block.level == 1:
                para(block.text, 14, _BLUE, "B", 7.5)
            else:
                para(block.text, 11, _NAVY, "B", 6)
            pdf.ln(0.8)
        elif isinstance(block, Para):
            para(block.text, 9.5 if block.muted else 10.5, _MUTED if block.muted else _TEXT)
            pdf.ln(1.6)
        elif isinstance(block, Quote):
            top = pdf.get_y()
            pdf.set_x(pdf.l_margin + 5)
            pdf.set_font(family, "", 10)
            pdf.set_text_color(*_TEXT)
            pdf.multi_cell(pdf.epw - 5, 5.2, text(f"\u201c{block.text}\u201d" if font else f'"{block.text}"'),
                           new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            if block.ref:
                pdf.set_x(pdf.l_margin + 5)
                pdf.set_font(family, "", 8.5)
                pdf.set_text_color(*_MUTED)
                pdf.multi_cell(pdf.epw - 5, 4.8, text(block.ref), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            bottom = pdf.get_y()
            if bottom > top:  # same page: draw the accent bar
                pdf.set_draw_color(*_BLUE)
                pdf.set_line_width(0.8)
                pdf.line(pdf.l_margin + 1.5, top + 0.5, pdf.l_margin + 1.5, bottom - 0.5)
                pdf.set_line_width(0.2)
            pdf.ln(2.5)
        elif isinstance(block, Table):
            pdf.set_font(family, "", 9)
            pdf.set_text_color(*_TEXT)
            pdf.set_draw_color(*_RULE)
            with pdf.table(col_widths=block.widths, text_align="LEFT", line_height=4.8, padding=1.6,
                           headings_style=FontFace(emphasis="BOLD", color=(255, 255, 255), fill_color=_BLUE),
                           cell_fill_color=_ZEBRA, cell_fill_mode=TableCellFillMode.ROWS,
                           borders_layout="HORIZONTAL_LINES") as table:
                head = table.row()
                for h in block.headers:
                    head.cell(text(h))
                for values in block.rows:
                    row = table.row()
                    for v in values:
                        row.cell(text(v))
            pdf.ln(3)
        elif isinstance(block, Lines):
            if not block.rows:
                para("No speech was transcribed.", 10, _MUTED)
                continue
            pdf.set_font(family, "", 9.5)
            pdf.set_text_color(*_TEXT)
            with pdf.table(col_widths=(17, 83), text_align="LEFT", line_height=4.9, padding=1.3,
                           first_row_as_headings=False, borders_layout="NONE",
                           cell_fill_color=_ZEBRA, cell_fill_mode=TableCellFillMode.ROWS) as table:
                for label, value in block.rows:
                    row = table.row()
                    row.cell(text(label), style=FontFace(color=_BLUE, size_pt=8.5))
                    row.cell(text(value))
            pdf.ln(3)
    return bytes(pdf.output())


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def render_docx(doc: Document) -> bytes:
    from docx import Document as WordDocument
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    navy, blue, muted, text_c = (RGBColor(*_NAVY), RGBColor(*_BLUE), RGBColor(*_MUTED), RGBColor(*_TEXT))
    word = WordDocument()
    section = word.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
    usable = section.page_width - section.left_margin - section.right_margin
    normal = word.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)

    def run(paragraph, value: str, *, size=None, bold=False, color=None):
        r = paragraph.add_run(_clean(value))
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
        return r

    header = section.header.paragraphs[0]
    run(header, f"{BRAND.upper()}  \u00b7  {TAGLINE}", size=8.5, bold=True, color=blue)
    run(header, f"   |   {doc.title}", size=8.5, color=muted)
    footer = section.footer.paragraphs[0]
    run(footer, f"Generated by {BRAND}  \u00b7  Page ", size=8, color=muted)
    _page_field(footer, OxmlElement, qn)

    def shade(cell, fill: str) -> None:
        props = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    for block in doc.blocks:
        if isinstance(block, Title):
            p = word.add_paragraph()
            run(p, block.text, size=22, bold=True, color=navy)
            p.paragraph_format.space_after = Pt(2)
            if block.subtitle:
                run(word.add_paragraph(), block.subtitle, size=11, color=muted)
        elif isinstance(block, Meta):
            table = word.add_table(rows=0, cols=2)
            for label, value in block.items:
                cells = table.add_row().cells
                run(cells[0].paragraphs[0], label.upper(), size=8.5, bold=True, color=muted)
                run(cells[1].paragraphs[0], value, size=9.5, color=text_c)
            for row in table.rows:
                row.cells[0].width = Cm(4.6)
                row.cells[1].width = usable - Cm(4.6)
            word.add_paragraph()
        elif isinstance(block, Heading):
            h = word.add_heading(level=1 if block.level == 1 else 2)
            run(h, block.text, color=blue if block.level == 1 else navy)
        elif isinstance(block, Para):
            p = word.add_paragraph()
            run(p, block.text, size=9.5 if block.muted else None, color=muted if block.muted else text_c)
        elif isinstance(block, Quote):
            p = word.add_paragraph(style="Quote")
            run(p, f"\u201c{block.text}\u201d", color=text_c)
            if block.ref:
                ref = word.add_paragraph()
                ref.paragraph_format.left_indent = Cm(0.6)
                run(ref, block.ref, size=8.5, color=muted)
        elif isinstance(block, Table):
            table = word.add_table(rows=1, cols=len(block.headers))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            total = sum(block.widths)
            for cell, h in zip(table.rows[0].cells, block.headers):
                run(cell.paragraphs[0], h, size=9, bold=True, color=RGBColor(255, 255, 255))
                shade(cell, "1D4ED8")
            for values in block.rows:
                cells = table.add_row().cells
                for cell, v in zip(cells, values):
                    run(cell.paragraphs[0], v, size=9, color=text_c)
            for row in table.rows:
                for cell, w in zip(row.cells, block.widths):
                    cell.width = int(usable * w / total)
            word.add_paragraph()
        elif isinstance(block, Lines):
            if not block.rows:
                run(word.add_paragraph(), "No speech was transcribed.", color=muted)
                continue
            for label, value in block.rows:
                p = word.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                run(p, label + "   ", size=8.5, bold=True, color=blue)
                run(p, value, color=text_c)

    word.core_properties.title = f"{doc.title} - {BRAND}"
    word.core_properties.author = BRAND
    buf = io.BytesIO()
    word.save(buf)
    return buf.getvalue()


def _page_field(paragraph, OxmlElement, qn) -> None:
    """Insert a PAGE field so Word shows the page number."""
    r = paragraph.add_run()
    for kind, value in (("begin", None), ("instr", " PAGE "), ("separate", None), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = value
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        r._r.append(el)
